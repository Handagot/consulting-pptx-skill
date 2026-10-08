"""Self-test for slides inserted into existing presentations (measure_deck / deck_pptx / check_deck --house / --xml-only).

The base "existing presentation" is generated on the fly from python-pptx's blank template (no pptx stored in repo).
Like test_checks.py, verifies one-by-one that uncorrected versions trigger FAIL while corrected versions pass.

  python3 -m unittest discover -s tests        # Skipped if python-pptx is not installed
"""
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
CHECK = ROOT / "scripts" / "check_deck.py"
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "examples"))

try:
    import pptx  # noqa: F401
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False


def make_house(path):
    """Generate 4 slides mimicking a house presentation: theme-default fonts, ruled tables, 14pt body, 20pt bold headings."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = 12192000, 6858000
    layout = next(l for l in prs.slide_layouts if l.name == "Title Only")
    title = layout.placeholders[0]   # Blank template is in 4:3 position, widen to 16:9
    title.left, title.top, title.width, title.height = Inches(0.5), Inches(0.3), Inches(12.33), Inches(0.9)
    for n in range(4):
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = f"既存の資料のページ {n + 1} は、罫線の表で書かれている"
        head = s.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(5), Inches(0.4))
        r = head.text_frame.paragraphs[0].add_run()
        r.text, r.font.size, r.font.bold = "見出し", Pt(20), True
        for i in range(3):
            body = s.shapes.add_textbox(Inches(0.5), Inches(2.2 + 0.6 * i), Inches(8), Inches(0.4))
            r = body.text_frame.paragraphs[0].add_run()
            r.text, r.font.size = "本文の行は 14pt で、書体を指定していない", Pt(14)
            if i:   # Differentiate 2 emphasis colors: more accent2 HEX (C0504D), fewer non-theme HEX (2E8B57)
                r = body.text_frame.paragraphs[0].add_run()
                r.text, r.font.size = ("強調の語" if i == 1 else "別の強調"), Pt(14)
                r.font.color.rgb = RGBColor(0xC0, 0x50, 0x4D) if i == 1 or n < 2 else RGBColor(0x2E, 0x8B, 0x57)
            ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(2.1 + 0.6 * i), Inches(12.8), Inches(2.1 + 0.6 * i))
            ln.line.color.rgb, ln.line.width = RGBColor(0x80, 0x80, 0x80), Pt(1.0 if i == 0 else 0.5)
    prs.save(str(path))


def check(path, *extra):
    r = subprocess.run([sys.executable, str(CHECK), str(path), *extra], capture_output=True, text=True)
    return r.returncode, [l for l in r.stdout.splitlines() if l.startswith("FAIL")]


def patched(src, dst, old, new, part="ppt/slides/slide1.xml"):
    """Create a copy of PPTX with only 1 modification (to construct uncorrected versions)."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part:
                assert old.encode("utf8") in data, old
                data = data.replace(old.encode("utf8"), new.encode("utf8"), 1)
            zout.writestr(item, data)
    return dst


@unittest.skipUnless(HAS_PPTX, "Executed when python-pptx is installed")
class HouseDeck(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from house_deck_example import build
        from measure_deck import measure

        cls.tmp = tempfile.TemporaryDirectory()
        cls.d = Path(cls.tmp.name)
        cls.house = cls.d / "house.pptx"
        make_house(cls.house)
        cls.skin = measure(cls.house)
        cls.skin_path = cls.d / "house.skin.json"
        cls.skin_path.write_text(json.dumps(cls.skin, ensure_ascii=False), encoding="utf8")
        cls.pages = cls.d / "pages.pptx"
        build(cls.house, cls.pages, cls.skin)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_measure_reads_the_house_style(self):
        k = self.skin
        self.assertEqual(k["layout"], "Title Only")
        self.assertTrue(k["fonts"]["inherit"])
        self.assertEqual((k["sizes"]["body"], k["sizes"]["head"]), (14.0, 20.0))
        # Rules are written in HEX (808080), but match 50% darker bg1 in theme, so stored as theme color
        self.assertEqual((k["rule"]["color"], k["rule"]["row"], k["rule"]["head"]), ("bg1|lumMod=50000", 0.5, 1.0))
        self.assertEqual(k["tables"]["native"], 0)
        self.assertEqual((k["margin"]["x0"], k["margin"]["x1"]), (0.5, 12.83))

    def test_pages_have_no_house_slides_and_pass(self):
        from pptx import Presentation
        self.assertEqual(len(Presentation(str(self.pages)).slides), 2)   # Base 4 slides are excluded
        code, fails = check(self.pages, "--house", str(self.skin_path))
        self.assertEqual((code, fails), (0, []))

    def test_pages_keep_fonts_and_tables_as_the_house_does(self):
        with zipfile.ZipFile(self.pages) as z:
            xml = "".join(z.read(n).decode("utf8") for n in z.namelist() if n.startswith("ppt/slides/slide"))
        self.assertNotIn("<a:latin", xml)   # Font is not set on run
        self.assertNotIn("<a:tbl>", xml)    # Table is textbox + rules
        self.assertNotIn('<a:srgbClr val="808080"/>', xml)   # Rule uses deck color as theme color
        self.assertIn('<a:schemeClr val="bg1"><a:lumMod val="50000"/></a:schemeClr>', xml)

    def test_colors_keep_several_options_and_follow_the_theme(self):
        k = self.skin
        self.assertEqual(k["color_options"]["emphasis"], ["accent2", "2E8B57"])   # Descending frequency; mapped to theme color
        self.assertEqual(k["colors"]["emphasis"], "accent2")
        self.assertEqual(k["colors"]["ng"], "FF0000")
        from deck_pptx import Deck
        d = Deck(self.house, self.skin)
        self.assertEqual((d.c("emphasis"), d.c("emphasis", 1), d.c("emphasis", 5)), ("accent2", "2E8B57", "accent2"))

    def test_to_theme_reads_tints(self):
        from measure_deck import to_theme
        theme = {"tx1": "000000", "bg1": "FFFFFF", "accent1": "4F81BD"}
        self.assertEqual(to_theme("4F81BD", theme), "accent1")
        self.assertEqual(to_theme("DCE6F2", theme), "accent1|lumMod=20000|lumOff=80000")   # "Lighter 80%"
        self.assertEqual(to_theme("123456", theme), "123456")                              # Non-theme colors kept as-is
        self.assertEqual(to_theme("accent3", theme), "accent3")

    def test_ng_color_comes_from_skin(self):
        from deck_pptx import Deck
        d = Deck(None, dict(self.skin, colors=dict(self.skin["colors"], ng="C00000")))
        s = d.slide("✕ の色は skin で変えられる")
        d.mark(s, "ng", 1, 2)
        xml = s._element.xml
        self.assertIn('val="C00000"', xml)
        self.assertNotIn('val="FF0000"', xml)

    def test_bad_theme_color_fires(self):
        bad = patched(self.pages, self.d / "bad_color.pptx", '<a:schemeClr val="accent1"', '<a:schemeClr val="FFC000"')
        code, fails = check(bad)
        self.assertEqual(code, 1)
        self.assertTrue(any("テーマ色の名前が不正" in f for f in fails), fails)

    def test_xml_only_stops_before_powerpoint(self):
        bad = patched(self.pages, self.d / "bad_color2.pptx", '<a:schemeClr val="accent1"', '<a:schemeClr val="FFC000"')
        self.assertEqual(check(bad, "--xml-only")[0], 1)
        self.assertEqual(check(self.pages, "--xml-only"), (0, []))

    def test_deck_pptx_refuses_bad_color(self):
        from deck_pptx import Deck
        d = Deck()
        s = d.slide("色の指定を誤ったら、ファイルを書く前に止まる")
        with self.assertRaises(ValueError):
            d.panel(s, 1, 2, 3, 1, fill="FFC000|lumMod=20000|lumOff")   # Missing conversion value
        with self.assertRaises(ValueError):
            d.panel(s, 1, 2, 3, 1, fill="accent9")                      # Non-existent theme color

    def test_explicit_font_fires_only_with_house(self):
        bad = patched(self.pages, self.d / "bad_font.pptx", 'sz="1400">', 'sz="1400"><a:latin typeface="Yu Gothic"/>')
        code, fails = check(bad, "--house", str(self.skin_path))
        self.assertTrue(any("書体を run に直指定" in f for f in fails), fails)
        self.assertFalse(any("書体" in f for f in check(bad)[1]))

    def test_title_outside_placeholder_fires(self):
        bad = patched(self.pages, self.d / "bad_title.pptx", '<p:ph type="title"/>', "")
        code, fails = check(bad, "--house", str(self.skin_path))
        self.assertTrue(any("タイトルがプレースホルダーに入っていない" in f for f in fails), fails)

    def test_dash_join_split_across_runs_is_detected(self):
        """Even if runs are split into 'Label', ' — ', 'Description', WARN dash concatenation at paragraph level."""
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation(str(self.house))
        s = prs.slides[0]
        tb = s.shapes.add_textbox(Inches(0.5), Inches(5), Inches(8), Inches(0.4))
        for piece in ("承認待ち", " — ", "平均3日"):
            r = tb.text_frame.paragraphs[0].add_run()
            r.text, r.font.size = piece, Pt(14)
        out = self.d / "dash_runs.pptx"
        prs.save(str(out))
        r = subprocess.run([sys.executable, str(CHECK), str(out)], capture_output=True, text=True)
        self.assertIn("ダッシュ", r.stdout)

    def test_potx_template_is_measured_and_built_on(self):
        """Verify measuring, building, and checking work even when house styles are distributed as .potx (PowerPoint template)."""
        from deck_pptx import Deck
        from measure_deck import measure
        potx = patched(self.house, self.d / "house.potx",
                       "presentationml.presentation.main+xml", "presentationml.template.main+xml",
                       part="[Content_Types].xml")
        skin = measure(potx)
        self.assertEqual((skin["layout"], skin["sizes"]["body"]), (self.skin["layout"], self.skin["sizes"]["body"]))
        d = Deck(potx, skin)
        d.slide("テンプレートの上に組んだページは、通常の pptx として保存される")
        out = self.d / "from_potx.pptx"
        d.save(out)
        with zipfile.ZipFile(out) as z:
            self.assertIn(b"presentationml.presentation.main+xml", z.read("[Content_Types].xml"))
        self.assertEqual(check(out, "--xml-only"), (0, []))
        self.assertEqual(check(potx, "--xml-only"), (0, []))

    def test_body_layout_is_found_in_a_one_of_each_sample(self):
        """Select body slide layout even in sample decks with 1 each of Title, Section, Body (avoid tie-break selecting Title)."""
        from pptx import Presentation
        from measure_deck import measure
        prs = Presentation()
        prs.slide_width, prs.slide_height = 12192000, 6858000
        for name in ("Title Slide", "Section Header", "Title and Content"):
            s = prs.slides.add_slide(next(l for l in prs.slide_layouts if l.name == name))
            s.shapes.title.text = f"{name} の見本"
        sample = self.d / "one_of_each.pptx"
        prs.save(str(sample))
        self.assertEqual(measure(sample)["layout"], "Title and Content")
        self.assertEqual(measure(sample, layout="Section Header")["layout"], "Section Header")   # Can also specify by name
        with self.assertRaises(SystemExit):
            measure(sample, layout="無いレイアウト")

    def test_line_break_in_shape_is_not_an_orphan(self):
        r = subprocess.run([sys.executable, str(CHECK), str(self.pages)], capture_output=True, text=True)
        self.assertNotIn("泣き別れ", r.stdout)   # Do not count soft break inside chevron shape as single line


if __name__ == "__main__":
    unittest.main()

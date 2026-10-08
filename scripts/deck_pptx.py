#!/usr/bin/env python3
"""Module for building editable slides directly onto the master of existing PowerPoint decks (slide-rules §8.7).

Converting HTML to PPTX (html_to_pptx.py) independently decides typefaces and margins, so inserting them into
an existing presentation can look inconsistent in typeface, font size, or table style.
When an existing target deck exists, build directly on top of that deck's master.

  from deck_pptx import Deck, P
  d = Deck("house.pptx", "house.skin.json")     # Base presentation and styles exported by measure_deck.py
  s = d.slide("Waiting time from application to payment is longest at approval", "Subtitle (omit if not needed)")
  d.head(s, d.x0, 1.6, 5.0, "Wait times by process stage")
  d.rule(s, 2.0, heavy=True)
  d.table(s, d.x0, 2.1, [3.0, 4.5, 4.6], ["Stage", "Current", "Target"], rows, [0.8] * len(rows))
  d.save("pages.pptx")                          # File containing only the spliced pages. Excludes base slides.

Rules to follow (align with base deck conventions; values are held in skin):
  - Do not set typeface, background color, or page numbers on the slide (if the deck delegates to the theme, don't set typeface on runs)
  - Place title and subtitle into layout placeholders
  - Build tables using text boxes and rules (bold header + rule underneath, thin rules between rows, no filled cells -- §6)
  - Pass colors as theme color references such as "accent2" or "accent2|lumMod=20000|lumOff=80000". HEX must be 6 digits
  - When the deck differentiates colors by role, select from candidates like d.c("emphasis", 1) (measure_deck color_options)
Dependency: python-pptx.
"""
import json
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure") and sys.stdout.encoding:
    sys.stdout.reconfigure(errors="replace")
if hasattr(sys.stderr, "reconfigure") and sys.stderr.encoding:
    sys.stderr.reconfigure(errors="replace")

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

EMU = 914400
RECT, OVAL = MSO_SHAPE.RECTANGLE, MSO_SHAPE.OVAL
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_SECTION_EXT = "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}"
SCHEME_COLORS = ("bg1", "tx1", "bg2", "tx2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
                 "hlink", "folHlink", "dk1", "lt1", "dk2", "lt2", "phClr")
COLOR_MODS = ("lumMod", "lumOff", "tint", "shade", "alpha", "satMod")

# Default style when no base deck is provided (testing with python-pptx blank template)
DEFAULT_SKIN = {
    "layout": "Title Only", "title": {"x": 0.5, "y": 0.3, "w": 12.33, "h": 0.9, "size": 28, "default_size": 44},
    "subtitle": None, "margin": {"x0": 0.5, "x1": 12.83, "body_top": 1.4, "bottom": 6.9},
    "sizes": {"head": 18, "body": 14, "dense": 12}, "fonts": {"inherit": True, "explicit": None},
    "textbox": {"insets": [0, 0, 0, 0], "lst_style": None, "bullet_level": None},
    "rule": {"color": "bg1|lumMod=50000", "row": 0.5, "head": 1.0}, "tables": {"native": 0},
    "colors": {"emphasis": "accent1", "panel": "accent1|lumMod=20000|lumOff=80000", "accent": "accent1", "ng": "FF0000"},
    "page_number": "none",
}


def E(v):
    return int(round(v * EMU))


def sub(parent, tag, **attrs):
    el = etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _shapes(o):
    """Return shapes container whether passed a Slide or GroupShape."""
    return o.shapes if hasattr(o, "shapes") else o


def color(parent, spec):
    """Add <a:solidFill>. spec is 'color|mod=val|...'. Color is 6-digit HEX or theme color name.
    Writing invalid HEX or unknown theme colors causes PowerPoint to fail opening the file (prompting repair), so stop here."""
    name, *mods = spec.split("|")
    fill = sub(parent, "a:solidFill")
    if re.fullmatch(r"[0-9A-Fa-f]{6}", name):
        c = sub(fill, "a:srgbClr", val=name.upper())
    elif name in SCHEME_COLORS:
        c = sub(fill, "a:schemeClr", val=name)
    else:
        raise ValueError(f"Invalid color specification: {spec!r} (must be 6-digit HEX or one of {SCHEME_COLORS})")
    for m in mods:
        k, _, v = m.partition("=")
        if k not in COLOR_MODS or not v.isdigit():
            raise ValueError(f"Invalid color modifier: {m!r} (must be one of {COLOR_MODS}=integer)")
        sub(c, "a:" + k, val=v)
    return fill


def P(text, sz=None, b=False, c=None, bullet=False, al="l", sb=None, sa=None):
    """A single paragraph. text is str or [(str, {sz,b,c}), ...]. \\n in str creates soft line breaks.
    Omit sz for body font size. bullet=True for bullets matching deck style. sb/sa are space before/after (pt)."""
    return dict(text=text, sz=sz, b=b, c=c, bullet=bullet, al=al, sb=sb, sa=sa)


class Deck:
    def __init__(self, template=None, skin=None, keep_slides=False):
        """template: Base deck (.pptx / .potx). skin: measure_deck.py output (path or dict). If omitted, measured on the fly.
        keep_slides=True retains base slides (use move_slides to position new slides)."""
        if template is None:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = 12192000, 6858000
            skin = skin or DEFAULT_SKIN
            t = DEFAULT_SKIN["title"]   # Widen blank template title placeholder to 16:9 width (default is 4:3 position)
            ph = next(l for l in self.prs.slide_layouts if l.name == "Title Only").placeholders[0]
            ph.left, ph.top, ph.width, ph.height = E(t["x"]), E(t["y"]), E(t["w"]), E(t["h"])
        else:
            from pptx_open import open_presentation
            self.prs = open_presentation(template)   # .potx (templates) can also serve as base. Saved as .pptx
            if skin is None:
                from measure_deck import measure
                skin = measure(template)
        if not isinstance(skin, dict):
            skin = json.loads(Path(skin).read_text(encoding="utf8"))
        self.skin = skin
        if not keep_slides:
            self._drop_slides()
        m, z = skin["margin"], skin["sizes"]
        self.x0, self.x1, self.top, self.bottom = m["x0"], m["x1"], m.get("body_top") or 1.4, m["bottom"]
        self.sz_head, self.sz_body, self.sz_dense = z["head"], z["body"], z["dense"]
        self.colors = dict(DEFAULT_SKIN["colors"], **{k: v for k, v in skin["colors"].items() if v})
        self.color_options = skin.get("color_options") or {}
        self.rule_color, self.rule_row, self.rule_head = skin["rule"]["color"], skin["rule"]["row"], skin["rule"]["head"]
        self.font = None if skin["fonts"]["inherit"] else skin["fonts"].get("explicit")
        self.insets = tuple(skin["textbox"]["insets"])
        self.lst_style = skin["textbox"].get("lst_style")
        self.bullet_level = skin["textbox"].get("bullet_level")

    def c(self, role, i=0):
        """Return i-th candidate color for role (emphasis / panel / accent). Falls back to default if index out of range."""
        opts = self.color_options.get(role) or []
        return opts[i] if i < len(opts) else self.colors[role]

    # ------------------------------------------------------------ Deck and Slides
    def _drop_slides(self):
        lst = self.prs.slides._sldIdLst
        for sld in list(lst):
            self.prs.part.drop_rel(sld.rId)
            lst.remove(sld)
        ext_lst = self.prs.part._element.find(qn("p:extLst"))
        for ext in (list(ext_lst) if ext_lst is not None else []):
            if ext.get("uri") == _SECTION_EXT:    # Drop sections list since it references removed slides
                ext_lst.remove(ext)

    def save(self, path):
        self.prs.save(str(path))

    def slide(self, title, subtitle=None, title_size=None, layout=None):
        """Add slide with title and subtitle in layout placeholders. Remove unused subtitle placeholder."""
        name = layout or self.skin.get("layout")
        lay = next((l for l in self.prs.slide_layouts if l.name.strip() == (name or "").strip()), None)
        if lay is None:
            raise KeyError(f"Layout '{name}' not found in base deck: {[l.name for l in self.prs.slide_layouts]}")
        s = self.prs.slides.add_slide(lay)
        t = self.skin.get("title") or {}
        size = title_size or t.get("size")
        done = False
        for ph in list(s.placeholders):
            kind = ph.placeholder_format.type
            if kind in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                self._fill_placeholder(ph, title, None if size == t.get("default_size") else size)
                done = True
            elif kind == PP_PLACEHOLDER.SUBTITLE:
                if subtitle:
                    self._fill_placeholder(ph, subtitle, None)
                else:
                    ph._element.getparent().remove(ph._element)
        if not done:   # Layout without title frame. check_deck interprets shapes whose name starts with Title as the title
            self.text(s, self.x0, 0.35, self.x1 - self.x0, 0.6, P(title, sz=size or 22, b=True), anchor="b", name="Title")
        return s

    @staticmethod
    def _fill_placeholder(ph, value, size):
        p = ph.text_frame.paragraphs[0]._p
        for child in list(p):
            p.remove(child)
        for i, part in enumerate(value.split("\n")):
            if i:
                sub(p, "a:br")
            r = sub(p, "a:r")
            rPr = sub(r, "a:rPr", lang="ja-JP", altLang="en-US")
            if size:
                rPr.set("sz", str(int(size * 100)))
            sub(r, "a:t").text = part

    def move_slides(self, indices, after):
        """Move slides at indices (0-indexed) immediately after slide 'after' (1-indexed). Also reorder in sections list."""
        lst = self.prs.slides._sldIdLst
        ids = list(lst)
        moving, anchor = [ids[i] for i in indices], ids[after - 1]
        for el in moving:
            lst.remove(el)
        pos = list(lst).index(anchor) + 1
        for k, el in enumerate(moving):
            lst.insert(pos + k, el)
        ext_lst = self.prs.part._element.find(qn("p:extLst"))
        tag = "{http://schemas.microsoft.com/office/powerpoint/2010/main}sldId"
        for sec in (ext_lst.iter(tag) if ext_lst is not None else []):
            if sec.get("id") == anchor.get("id"):
                parent, at = sec.getparent(), list(sec.getparent()).index(sec) + 1
                for k, el in enumerate(moving):
                    new = etree.Element(tag)
                    new.set("id", el.get("id"))
                    parent.insert(at + k, new)
                break

    # ------------------------------------------------------------ Text
    def bullets(self, items, sz=None, **kw):
        return [P(t, sz=sz or self.sz_dense, bullet=True, **kw) for t in items]

    def _rpr(self, parent, sz, b, c):
        rPr = sub(parent, "a:rPr", lang="ja-JP", altLang="en-US", sz=str(int(round(sz * 100))))
        if b:
            rPr.set("b", "1")
        if c:
            color(rPr, c)
        if self.font:   # Set only when base deck explicitly specifies typeface on runs
            sub(rPr, "a:latin", typeface=self.font)
            sub(rPr, "a:ea", typeface=self.font)

    def _para(self, p, s, in_shape):
        pPr = sub(p, "a:pPr")
        styled = s["bullet"] and self.lst_style and self.bullet_level is not None and not in_shape
        if styled:
            pPr.set("lvl", str(self.bullet_level))
        elif s["bullet"]:
            pPr.set("marL", str(E(0.2)))
            pPr.set("indent", str(-E(0.2)))
        if s["al"] != "l" or in_shape:
            pPr.set("algn", {"l": "l", "c": "ctr", "r": "r"}[s["al"]])
        sb, sa = s["sb"], s["sa"]
        if in_shape or not self.lst_style:   # Textboxes without list styles manage space before/after explicitly
            sb = 0 if sb is None else sb
            sa = (3 if s["bullet"] else 0) if sa is None else sa
        if sb is not None:
            sub(sub(pPr, "a:spcBef"), "a:spcPts", val=int(sb * 100))
        if sa is not None:
            sub(sub(pPr, "a:spcAft"), "a:spcPts", val=int(sa * 100))
        if s["bullet"] and not styled:   # Bullet markers are styled, not typed as text characters (§7.3)
            sub(pPr, "a:buFont", typeface="Arial")
            sub(pPr, "a:buChar", char="•")
        elif not s["bullet"] and not self.lst_style:
            sub(pPr, "a:buNone")
        runs = s["text"] if isinstance(s["text"], list) else [(s["text"], {})]
        base = s["sz"] or self.sz_body
        last = base
        for txt, ov in runs:
            sz, b, c = ov.get("sz", base), ov.get("b", s["b"]), ov.get("c", s["c"])
            last = sz
            for i, part in enumerate(txt.split("\n")):
                if i:
                    self._rpr(sub(p, "a:br"), sz, b, c)
                if part:
                    r = sub(p, "a:r")
                    self._rpr(r, sz, b, c)
                    sub(r, "a:t").text = part
        sub(p, "a:endParaRPr", lang="en-US", sz=str(int(round(last * 100))))

    def _body(self, txBody, paras, anchor, insets, in_shape):
        for child in list(txBody):
            txBody.remove(child)
        l, t, r, b = insets
        bodyPr = sub(txBody, "a:bodyPr", vert="horz", wrap="square", lIns=E(l), tIns=E(t), rIns=E(r), bIns=E(b),
                     rtlCol="0", anchor={"t": "t", "m": "ctr", "b": "b"}[anchor])
        sub(bodyPr, "a:noAutofit")
        if in_shape or not self.lst_style:
            sub(txBody, "a:lstStyle")
        else:   # Mirror list styling from deck textboxes (preserves bullet indent, marker, and space after)
            txBody.append(etree.fromstring(self.lst_style.replace("<a:lstStyle", f'<a:lstStyle xmlns:a="{_A}"', 1)))
        for s in ([paras] if isinstance(paras, dict) else paras):
            self._para(sub(txBody, "a:p"), s, in_shape)

    def text(self, sl, x, y, w, h, paras, anchor="t", name=None):
        """Textbox built to match deck conventions (inherits padding, list styles; no autofit)."""
        t = _shapes(sl).add_textbox(E(x), E(y), E(w), E(h))
        self._body(t._element.txBody, paras, anchor, self.insets, False)
        if name:
            t.name = name
        return t

    # ------------------------------------------------------------ Shapes
    @staticmethod
    def _paint(s, fill=None, stroke=None, lw=0.5, alpha=None):
        spPr = s._element.spPr
        for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:ln"):
            for el in spPr.findall(qn(tag)):
                spPr.remove(el)
        if fill:
            f = color(spPr, fill)
            if alpha is not None:
                sub(f[0], "a:alpha", val=int(alpha * 100000))
        else:
            sub(spPr, "a:noFill")
        ln = sub(spPr, "a:ln", w=int(lw * 12700))
        if stroke:
            color(ln, stroke)
        else:
            sub(ln, "a:noFill")
        st = s._element.find(qn("p:style"))   # Prevent inheriting theme defaults (shadow, borders)
        if st is not None:
            s._element.remove(st)

    def box(self, sl, x, y, w, h, fill=None, stroke=None, lw=0.5, kind=RECT, alpha=None, adj=None, rot=None, name=None):
        """Shape. Use either with fill and no border, or border and no fill (§5.3)."""
        s = _shapes(sl).add_shape(kind, E(x), E(y), E(w), E(h))
        self._paint(s, fill, stroke, lw, alpha)
        if adj is not None:
            for i, v in enumerate(adj if isinstance(adj, (list, tuple)) else [adj]):
                s.adjustments[i] = v
        if rot:
            s.rotation = rot
        if name:
            s.name = name
        return s

    def tbox(self, sl, x, y, w, h, paras, fill=None, anchor="m", kind=RECT, adj=None, name=None,
             pad=(0.1, 0.05, 0.1, 0.05), alpha=None):
        """Filled shape with text (banner, card, chevron, circle)."""
        s = self.box(sl, x, y, w, h, fill=fill, kind=kind, adj=adj, name=name, alpha=alpha)
        txBody = s._element.find(qn("p:txBody"))
        if txBody is None:
            txBody = sub(s._element, "p:txBody")
        self._body(txBody, paras, anchor, pad, True)
        return s

    def line(self, sl, x1, y1, x2, y2, stroke=None, lw=None, arrow=False):
        c = _shapes(sl).add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        self._paint(c, None, stroke or self.rule_color, lw or self.rule_row)
        spPr = c._element.spPr
        spPr.remove(spPr.find(qn("a:noFill")))
        if arrow:
            sub(spPr.find(qn("a:ln")), "a:tailEnd", type="triangle", w="med", len="med")
        return c

    def poly(self, sl, pts, stroke=None, lw=1.0, close=False, fill=None):
        fb = _shapes(sl).build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
        fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=close)
        s = fb.convert_to_shape()
        self._paint(s, fill, stroke, lw)
        return s

    # ------------------------------------------------------------ Combinations matching deck conventions
    def rule(self, sl, y, x0=None, x1=None, heavy=False):
        """Horizontal rule. heavy=True under headings; thin between rows. Color and width match deck."""
        return self.line(sl, self.x0 if x0 is None else x0, y, self.x1 if x1 is None else x1, y,
                         lw=self.rule_head if heavy else self.rule_row)

    def head(self, sl, x, y, w, label, sz=None, h=0.36):
        """Section heading (bold, bottom-aligned). Follow immediately with rule(..., heavy=True)."""
        return self.text(sl, x, y, w, h, P(label, sz=sz or self.sz_head, b=True), anchor="b")

    def panel(self, sl, x, y, w, h, fill=None, name="Emphasis Panel"):
        """Panel placed behind highlighted column/row. Place before text (renders in background)."""
        return self.box(sl, x, y, w, h, fill=fill or self.colors["panel"], name=name)

    def num_circle(self, sl, x, y, n, d=0.31, fill=None):
        return self.tbox(sl, x, y, d, d, P(str(n), c="bg1", al="c"), fill=fill or self.colors["accent"], kind=OVAL,
                         pad=(0, 0, 0, 0), name=f"Number {n}")

    def triangle(self, sl, x, y, w, h, direction="down", fill=None):
        """Filled triangle showing process flow (§4.27). direction is down or right."""
        fill = fill or self.colors["accent"]
        if direction == "right":   # Rotate box swaps width and height
            cx, cy = x + w / 2, y + h / 2
            return self.box(sl, cx - h / 2, cy - w / 2, h, w, fill=fill, kind=MSO_SHAPE.ISOSCELES_TRIANGLE, rot=90)
        return self.box(sl, x, y, w, h, fill=fill, kind=MSO_SHAPE.ISOSCELES_TRIANGLE, rot=180)

    def chevrons(self, sl, x, y, w, h, items, fills=None, depth=0.24, gap=0.05, pad_left=0.12):
        """Chevron sequence (§4.18). items is sequence of paragraphs (P or list of P). Returns list of (left, width)."""
        n = len(items)
        pitch = (w + gap) / n
        out = []
        for i, paras in enumerate(items):
            sx, last = x + i * pitch, i == n - 1
            self.tbox(sl, sx, y, pitch - gap if last else pitch - gap + depth, h, paras,
                      fill=(fills or [self.colors["accent"]] * n)[i],
                      kind=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, adj=depth / h,
                      pad=(pad_left, 0, 0.02, 0), name=f"Chevron {i + 1}")
            out.append((sx, pitch - gap))
        return out

    def mark(self, sl, kind, x, y, s=0.2):
        """Semantic checkmarks ✓, ✕, △ instead of ○, ✕, △. Drawn as shapes, not text characters (§4.27). kind is ok / ng / tri."""
        g = _shapes(sl).add_group_shape()
        if kind == "ok":
            self.poly(g, [(x + 0.08 * s, y + 0.52 * s), (x + 0.38 * s, y + 0.84 * s), (x + 0.94 * s, y + 0.16 * s)],
                       self.colors["emphasis"], 2.25)
        elif kind == "ng":   # Semantic red is never aligned to brand colors (§5.8)
            self.line(g, x + 0.14 * s, y + 0.14 * s, x + 0.86 * s, y + 0.86 * s, self.colors["ng"], 2.0)
            self.line(g, x + 0.86 * s, y + 0.14 * s, x + 0.14 * s, y + 0.86 * s, self.colors["ng"], 2.0)
        else:
            self.box(g, x + 0.08 * s, y + 0.12 * s, 0.84 * s, 0.76 * s, stroke=self.rule_color, lw=1.5,
                     kind=MSO_SHAPE.ISOSCELES_TRIANGLE)
        g.name = {"ok": "mark ✓", "ng": "mark ✕"}.get(kind, "mark △")
        return g

    def table(self, sl, x, y, col_w, headers, rows, row_h, head_h=0.3, gap=0.12, sz=None, first_bold=True,
              pad_top=0.1, x1=None, anchor="t"):
        """Build structured table using text boxes and rules (cells are not filled; §6).

        headers: Column headers (None for no header row). rows: Row cell data. Cells are str / P / list of P /
        None (None reserved for caller to place circles or shapes). row_h is list of row heights. anchor="m" centers vertically.
        Returns {"xs": col lefts, "ys": row tops, "bottom": table bottom}. Highlighted columns/rows should layer panel() beforehand.
        """
        sz = sz or self.sz_body
        xs, cx = [], x
        for wdt in col_w:
            xs.append(cx)
            cx += wdt
        right = x1 if x1 is not None else cx - gap
        top = y
        if headers:
            for hx, wdt, label in zip(xs, col_w, headers):
                if label:
                    self.text(sl, hx, y, wdt - gap, head_h, P(label, sz=sz, b=True), anchor="b")
            top = y + head_h + 0.05
            self.rule(sl, top, x, right, heavy=True)
        ys = []
        for ri, (cells, rh) in enumerate(zip(rows, row_h)):
            if ri:
                self.rule(sl, top, x, right)   # Between rows only. Do not draw rule under the final row (§5.4)
            ys.append(top)
            for ci, (cell, hx, wdt) in enumerate(zip(cells, xs, col_w)):
                if cell is None:
                    continue
                paras = P(cell, sz=sz, b=(first_bold and ci == 0)) if isinstance(cell, str) else cell
                if anchor == "m":
                    self.text(sl, hx, top, wdt - gap, rh, paras, anchor="m")
                else:
                    self.text(sl, hx, top + pad_top, wdt - gap, rh - pad_top, paras)
            top += rh
        return dict(xs=xs, ys=ys, bottom=top)

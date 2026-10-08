#!/usr/bin/env python3
"""Automated slide design rule checker (unified PPTX / HTML).

Usage:
  python3 check_deck.py out.pptx
  python3 check_deck.py deck.html
  python3 check_deck.py assets/SuperTemplate_62type.pptx --template   # When inspecting template library itself
  python3 check_deck.py deck.html --forbid ~/.config/deck-forbidden-terms.txt   # FAIL on remaining client names/jargon
  python3 check_deck.py pages.pptx --house house.skin.json   # Pages inserted into house deck (slide-rules §8.7)
  python3 check_deck.py out.pptx --xml-only                  # Check only XML errors preventing PowerPoint from opening

Canonical Rules: references/slide-rules.md
Exit code: 1 if any FAIL occurs.
"""
import re
import sys
import zipfile
from pathlib import Path

if hasattr(sys.stdout, "reconfigure") and sys.stdout.encoding:
    sys.stdout.reconfigure(errors="replace")
if hasattr(sys.stderr, "reconfigure") and sys.stderr.encoding:
    sys.stderr.reconfigure(errors="replace")

# Prohibited legacy colors for custom branding (6-digit HEX, without #)
OLD_COLORS = []
TITLE_MAX = 40
EMU_W, EMU_H = 12192000, 6858000

fails, warns = [], []


def fail(msg):
    fails.append(msg)


def warn(msg):
    warns.append(msg)


# --- XML errors preventing PowerPoint from opening (slide-rules §8.7) -----------------
# Presentations with HEX or invalid names in theme color slots cause PowerPoint to demand repair.
# Stopping before opening avoids leaving persistent repair dialogs on the user's machine.
SCHEME_COLORS = {"bg1", "tx1", "bg2", "tx2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
                 "hlink", "folHlink", "dk1", "lt1", "dk2", "lt2", "phClr"}
XML_ONLY = False   # --xml-only: Stop after checking XML validity
HOUSE = None       # --house: Style measured by measure_deck.py for insertion


def check_xml_validity(idx, xml):
    bad = sorted(set(re.findall(r'<a:schemeClr val="([^"]*)"', xml)) - SCHEME_COLORS)
    if bad:
        fail(f"p{idx}: Invalid theme color name / テーマ色の名前が不正 {bad} (Write HEX in srgbClr; PowerPoint cannot open -- §8.7)")
    if re.search(r'<a:ext cx="-\d+"|<a:ext cx="\d+" cy="-\d+"', xml):
        fail(f"p{idx}: Negative shape dimensions / 図形の大きさが負の値 (PowerPoint cannot open -- §8.7)")
    szs = [int(v) for v in re.findall(r'<a:(?:rPr|endParaRPr|defRPr)[^>]*\bsz="(\d+)"', xml)]
    if any(not 100 <= v <= 400000 for v in szs):
        fail(f"p{idx}: Font size sz out of range / 文字の大きさ sz が範囲外 (100–400000; PowerPoint cannot open -- §8.7)")


def check_house(idx, slide, xml):
    """Verifies alignment with house deck master style (slide-rules §8.7). HOUSE is from measure_deck.py."""
    if HOUSE.get("title") and slide.shapes.title is None:
        fail(f"p{idx}: Title not in layout placeholder / タイトルがプレースホルダーに入っていない (Position and font are governed by house layout -- §8.7)")
    faces = sorted(set(re.findall(r'<a:(?:latin|ea) typeface="((?!\+)[^"]*)"', xml)))
    if faces and HOUSE["fonts"]["inherit"]:
        fail(f"p{idx}: Explicit font override on run / 書体を run に直指定 {faces[:4]} (House deck relies on theme; overrides will drift -- §8.7)")
    elif faces and set(faces) - {HOUSE["fonts"].get("explicit")}:
        warn(f"p{idx}: Font differs from house deck / 資料と違う書体 {sorted(set(faces) - {HOUSE['fonts'].get('explicit')})[:4]} (House font is {HOUSE['fonts'].get('explicit')})")
    if "<a:tbl>" in xml and not HOUSE["tables"]["native"]:
        warn(f"p{idx}: Native PowerPoint table object used / PowerPoint の表オブジェクト (House deck uses textboxes + rules; use deck_pptx table -- §8.7)")
    if 'type="slidenum"' in xml and HOUSE.get("page_number") == "layout":
        warn(f"p{idx}: Slide number placed directly on slide (House layout outputs slide number automatically -- §8.7)")
    used = set(v.upper() for v in re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', xml)) | set(re.findall(r'<a:schemeClr val="([^"]+)"', xml))
    ng = (HOUSE.get("colors") or {}).get("ng") or "FF0000"
    extra = sorted(used - set(HOUSE.get("palette") or used) - {ng.split("|")[0].upper(), "FF0000", "FFFFFF", "000000", "bg1", "tx1"})
    if extra:
        warn(f"p{idx}: Color not in house palette / 資料に無い色 {extra[:6]} (Select from house theme colors -- §8.7)")


# True only when inspecting template libraries directly (--template).
# In production decks, remaining placeholders like "Text N" or "Label N" trigger FAIL.
TEMPLATE_MODE = False
STRICT_LEN = True

# Representative terminology variation pairs (slide-rules §7.6: one term per deck). WARN if both appear.
PART_TYPE_NAMES = set()

TERM_VARIANTS = [
    ("メモリー", r"メモリー", "メモリ", r"メモリ(?!ー)(?!・CPU)"),
    ("紐付", r"紐付", "紐づ", r"紐づ"),
    ("ユーザー", r"ユーザー", "利用者", r"利用者"),
    ("アセット", r"アセット", "資産", r"資産"),
    ("フォルダー", r"フォルダー", "フォルダ", r"フォルダ(?!ー)"),
    ("サーバー", r"サーバー", "サーバ", r"サーバ(?!ー)"),
    ("メンバー", r"メンバー", "メンバ", r"メンバ(?!ー)"),
    ("コンピューター", r"コンピューター", "コンピュータ", r"コンピュータ(?!ー)"),
    ("問い合わせ", r"問い合わせ", "問合せ", r"問合せ"),
    ("adviser", r"\badviser\b", "advisor", r"\badvisor\b"),
    ("behavior", r"\bbehavior\b", "behaviour", r"\bbehaviour\b"),
]

# AI smell words and phrases (slide-rules §7.9 / references/ai-smell-lexicon.md). High-confidence WARN.
AI_SMELL_WORDS = [
    "まさに", "非常に", "極めて", "圧倒的", "画期的", "革新的", "次世代の",
    "過言ではありません", "に他なりません", "シームレス", "シナジー", "ソリューション",
    "エンドツーエンド", "ブラッシュアップ", "付加価値",
    "寄り添い", "伴走し", "二人三脚", "さらなる高みへ", "邁進",
    "昨今", "変化の激しい", "という点において", "の観点から",
    "させていただきます", "いただけますと幸いです",
    "と言えるでしょう", "と考えられます", "することが可能です",
    "詳細は別途", "別途ご説明", "追ってご連絡", "今後検討してまいります", "詳細は後日",
    "game-changer", "groundbreaking", "next-generation", "paradigm shift",
    "synergy", "seamlessly", "holistic approach", "actionable insights",
    "details will be provided separately", "to be determined at a later stage",
]


def check_production_meta(slides):
    """Verifies that production metadata (skill name, repository) is not visible on content slides.
    Attribution may appear ONLY on the final slide's source line. Anywhere else triggers FAIL.
    """
    def visible(s):
        s = re.sub(r"<(script|style)[^>]*>.*?</\1>|<!--.*?-->", " ", s, flags=re.S)
        return re.sub(r"<[^>]+>", " ", s)
    if TEMPLATE_MODE:
        return
    for i, s in enumerate(slides, 1):
        if i == len(slides):
            continue
        v = visible(s)
        m = re.search(r"consulting-pptx-skill", v)
        if m:
            fail(f"p{i}: Production metadata displayed on slide / 制作メタ: 「{m.group(0)}」 (Tool name may appear only in final slide source line)")
            continue
        m = re.search(r"(?:本資料は[^。<]{0,40}で作成|Created with [^.<]{0,40})", v)
        if m:
            warn(f"p{i}: 「{m.group(0)}」 (Move production credits to final slide source line; permitted if disclaimer / で作成)")
        elif re.search(r"github\.com", v):
            warn(f"p{i}: github.com displayed on slide (remove if production metadata; permitted if citing public repository)")


def pptx_text(xml):
    """Returns slide XML visible text as one line per paragraph (<a:p>)."""
    paras = re.findall(r"<a:p\b.*?</a:p>", xml, re.S)
    if not paras:
        return "\n".join(re.findall(r"<a:t>(.*?)</a:t>", xml, re.S))
    return "\n".join("".join(re.findall(r"<a:t>(.*?)</a:t>", p, re.S)) for p in paras)


def check_ai_smell(pages):
    """slide-rules §7.9: Detects high-confidence AI smell words (WARN)."""
    hits = {}
    for i, txt in pages:
        found = [w for w in AI_SMELL_WORDS if w.lower() in txt.lower()]
        if found:
            hits[i] = found
    if hits:
        detail = ", ".join(f"p{i}: 「{'/'.join(ws[:3])}」" for i, ws in sorted(hits.items())[:6])
        warn(f"AI smell vocabulary detected: {detail} (§7.9 / ai-smell-lexicon.md; replace with plain, active verbs)")

    def _dash_joined(txt):
        for line in txt.split("\n"):
            t = line.strip()
            if "—" not in t or re.fullmatch(r"—+(?:\s*[（(][^）)]*[）)])?", t):
                continue
            return True
        return False
    dash_pages = sorted({i for i, txt in pages if _dash_joined(txt)})
    if dash_pages:
        warn(f"Dash concatenation / ダッシュ p{dash_pages} (Classic AI style; replace with period, colon, or parenthesis. §7.9)")


def check_terms(pages):
    """pages: [(idx, text), ...] WARN if both variations appear in the same deck (§7.6)."""
    for la, pa, lb, pb in TERM_VARIANTS:
        hits_a = sorted({i for i, t in pages if re.search(pa, t, re.I)})
        hits_b = sorted({i for i, t in pages if re.search(pb, t, re.I)})
        if hits_a and hits_b:
            warn(f"Potential terminology inconsistency: 「{la}」p{hits_a} and 「{lb}」p{hits_b} both appear / 表記ゆれ疑い (§7.6)")


# --- Numerical consistency (slide-rules §7.6) ----------------------------------------
NUM_UNITS = (r"円|ドル|ユーロ|人|名|件|社|団体|店舗|拠点|校|戸|世帯|台|個|本|枚|冊|回|倍|%|％|pt|ポイント|"
             r"時間|日|か月|ヶ月|週|kWh|MWh|GWh|kW|MW|GW|kg|km|m3|㎥|m2|㎡|ha|ヘクタール|g|t|トン|m|L|"
             r"USD|EUR|GBP|JPY|users|clients|accounts|sites|stores|hours|days|weeks|months|years")
NUM_LABEL = re.compile(
    r"(?:(?P<year>(?:19|20)\d{2})(?:年度?|s)?(?:の|時点の|末の|末時点の)?)?"
    r"(?P<label>[一-龥々ァ-ヶーA-Za-z]{2,20})(?:は|が|の|：|:|＝|=)?\s*(?:約|およそ|計|合計|approx\.?)?\s*"
    r"(?P<num>\d[\d,]*(?:\.\d+)?)\s*(?P<unit>万|億|兆|千|k|M|B|T)?(?P<base>" + NUM_UNITS + r")(?![A-Za-z])",
    re.I
)
_ZEN = str.maketrans("０１２３４５６７８９，．％：＝", "0123456789,.%:=")
_MULT = {"千": 1e3, "万": 1e4, "億": 1e8, "兆": 1e12, "k": 1e3, "M": 1e6, "B": 1e9, "T": 1e12}
_UNIT_ALIAS = {"名": "人", "％": "%", "ポイント": "pt", "ヶ月": "か月", "㎥": "m3", "㎡": "m2",
               "ヘクタール": "ha", "トン": "t"}
_GENERIC_LABELS = {"以上", "以下", "未満", "超", "平均", "合計", "全体", "最大", "最小", "最高", "最低",
                   "前年", "前年比", "同期", "うち", "残り", "目標", "実績", "約", "計",
                   "total", "average", "target", "actual", "approx", "max", "min"}


def _num_facts(text):
    out = []
    for m in NUM_LABEL.finditer(text.translate(_ZEN)):
        if m["label"].lower() in _GENERIC_LABELS:
            continue
        label = (m["year"] + "年:" if m["year"] else "") + m["label"]
        base = _UNIT_ALIAS.get(m["base"], m["base"])
        try:
            val = round(float(m["num"].replace(",", "")) * _MULT.get(m["unit"] or "", 1), 6)
        except ValueError:
            continue
        out.append(((label, base), val, m.group(0).strip()))
    return out


def check_number_consistency(pages):
    seen = {}
    for i, t in pages:
        for key, val, raw in _num_facts(t):
            seen.setdefault(key, []).append((i, val, raw))
    for (label, base), hits in seen.items():
        vals = {v for _, v, _ in hits}
        pages_hit = {i for i, _, _ in hits}
        if len(vals) > 1 and len(pages_hit) > 1:
            detail = " / ".join(f"p{i}「{raw}」" for i, _, raw in hits[:4])
            warn(f"Numerical inconsistency / 数値の平仄疑い: 「{label.replace(':', '')}」 differs across {detail} (§7.6)")


# --- Body placeholders (slide-rules §2.8 / README) -----------------------------------
BODY_PLACEHOLDER = re.compile(
    r"Text\s*\d+|ラベル\s*\d+|Label\s*\d+|タイトル\s*\d+|Title\s*\d+|Source\s*\d+|YYYY|パーツ\s*\d+\s*[｜|]|Part\s*\d+\s*[|｜]|ダミー|Dummy|^会社名$|^連絡先$|^Company Name$|^Contact$")


def _part_type_names():
    names = set()
    root = Path(__file__).resolve().parent.parent / "templates"
    for f in ("freeform_parts_16x9.html", "freeform_parts_more_16x9.html"):
        fp = root / f
        if not fp.exists():
            continue
        for m in re.finditer(r"(?:パーツ|Part)\s*\d+\s*[｜|]\s*([^<]+)<", fp.read_text(encoding="utf8", errors="ignore")):
            names.add(m.group(1).strip())
    names.update({
        "表紙", "全体のマップ", "目次", "章扉", "矢羽（プロセス・変遷）", "前提→帰結の2カラム",
        "大型数値の表＋読み取り", "カード 2×2", "軸のある表", "裏表紙", "主張パネル＋図",
        "打ち手の効果表", "状態ヒートマップ＋右コメント", "充足度評価表（ハーベイボール）",
        "割合のドットマトリクス", "進捗バブル行列", "分布の順位棒＋注記", "注記つき散布図",
        "柱＋土台", "対向シェブロン", "外部動向の根拠グリッド", "比例円の対比",
        "増減の縦棒＋左右合計", "シナリオ線＋成長率チップ", "調査の土台", "課題と打ち手の2カラム",
        "セパレーター（章扉＝アジェンダ再掲）",
        "Title Page", "Overview Map", "Table of Contents", "Section Divider", "Chevron Steps",
        "Premise -> Conclusion 2-Column", "Stat Table + Readout", "2x2 Card Grid", "Table with Axes",
        "Back Cover", "Claim Panel + Visual", "Initiative Impact Table", "Status Heatmap + Right Commentary",
        "Evaluation Table (Harvey Balls)", "Proportion Dot Matrix", "Progress Bubble Matrix",
        "Ranked Distribution Bars + Notes", "Annotated Scatter Plot", "Pillars + Foundation",
        "Opposing Chevrons", "External Evidence Grid", "Proportional Circle Comparison",
        "Variance Vertical Bars + Net Total", "Scenario Lines + CAGR Chips", "Research Foundation",
        "Problem & Countermeasure 2-Column", "Separator (Chapter Divider = Repeated Agenda)",
        "エグゼクティブサマリー", "調査の根拠", "2つの大型数値", "KPIダッシュボード",
        "単一チャート＋示唆", "積み上げ棒", "ウォーターフォール", "完全版ウォーターフォール",
        "スモールマルチプル", "比較表", "シナリオ比較表", "リスク一覧表", "横軸評価表",
        "ヒートマップ表", "2×2マトリクス", "プロセス×観点の行列", "入れ子行の行列",
        "時系列マトリクス", "テーマカード", "提言の柱", "番号つき打ち手",
        "Situation・Complication・Resolution", "課題と打ち手の対応", "課題→原因→解決",
        "イシューツリー", "現状と目指す姿", "計算ロジックの流れ", "プロセスの段階",
        "循環サイクル", "矢羽の段階（全体像の見出し）", "バリューチェーン", "分岐と判断",
        "ロードマップ", "ガントチャート", "意思決定ページ",
        "Executive Summary", "Big Stat Comparison", "KPI Dashboard",
        "Single Chart + Strategic Implications", "Stacked Bar Breakdown", "Contribution Bridge / Waterfall",
        "Full Variance Bridge", "Small Multiples Comparison", "Multi-Option Evaluation Matrix",
        "Scenario Comparison Table", "Risk Assessment & Mitigation Matrix", "Horizontal Axis Evaluation Table",
        "Heatmap Matrix", "2x2 Strategic Matrix", "Process x Perspective Matrix",
        "Nested Row Hierarchy Matrix", "Timeline Matrix", "Theme Card Grid", "Recommendation Pillars",
        "Numbered Strategic Imperatives", "Situation, Complication, Resolution",
        "Issue-to-Solution Mapping", "Issue -> Root Cause -> Solution", "Issue Tree",
        "Current vs. Target State", "Calculation Logic Flow", "Process Flow Stages",
        "Closed-Loop Cycle", "Chevron Rail Stages", "Chevron Rail Stages (Top-Level Headings)",
        "Value Chain Chevrons", "Decision Fork", "Implementation Roadmap", "Gantt Chart",
        "Executive Decision Page",
    })
    return names


def check_body_placeholders(idx, leaf_texts, title):
    if TEMPLATE_MODE:
        return
    hits = sorted({t for t in leaf_texts if BODY_PLACEHOLDER.search(t)})
    if hits:
        fail(f"p{idx}: Template placeholder remains in body / 本文にテンプレのプレースホルダー ×{len(hits)}: {' / '.join(h[:20] for h in hits[:4])}")
    if title and title.strip() in PART_TYPE_NAMES:
        fail(f"p{idx}: Title remains raw archetype name / 型名のまま 「{title.strip()}」 (Rewrite as assertive takeaway -- §2.8)")


# --- Single sentence per text block (slide-rules §7.24) -----------------------------
SENTENCE_END = re.compile(r"[。！？!?](?=\s*\S)")


def check_multi_sentence(idx, leaf_texts):
    bad = []
    for t in leaf_texts:
        if t.startswith(("出典", "注", "※", "Source:", "Note:", "Source", "Note")):
            continue
        core = re.sub(r"（[^）]*）|\([^)]*\)|「[^」]*」|“[^”]*”|\"[^\"]*\"", "", t)
        if len(SENTENCE_END.findall(core.strip())) >= 1:
            bad.append(t)
    if bad:
        fail(f"p{idx}: Multiple sentences in single text block / 2文以上 ×{len(bad)}: 「{bad[0][:36]}…」 (Decompose into 1 item per sentence)")


# --- Forbidden terms (Client names, internal codes) ----------------------------------
FORBIDDEN_TERMS = []


def load_forbidden(path):
    terms = []
    for line in Path(path).expanduser().read_text(encoding="utf8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        terms.append(re.compile(line[3:]) if line.startswith("re:") else re.compile(re.escape(line)))
    return terms


def check_forbidden(pages, raw=""):
    if not FORBIDDEN_TERMS:
        return
    for pat in FORBIDDEN_TERMS:
        hit_pages = sorted({i for i, t in pages if pat.search(t)})
        in_markup = bool(raw and pat.search(re.sub(r">[^<]*<", "><", raw)))
        if hit_pages or in_markup:
            where = f"p{hit_pages}" if hit_pages else "HTML comments/attributes (コメント/属性)"
            fail(f"Forbidden term / 禁止語 list match found in {where} (Row {FORBIDDEN_TERMS.index(pat)+1} in list)")


def check_title(idx, title, explicit_break=False):
    t = title.strip()
    if not t:
        warn(f"p{idx}: Empty title / タイトルが空 (Permitted for title page / divider)")
        return
    if re.search(r"(です|ます|でした|ました)[。．.]?$", t):
        fail(f"p{idx}: Title ends in polite conversational phrasing / ですます調: 「{t}」")
    tlen = sum(0.5 if ord(ch) < 0x3000 else 1 for ch in t)
    if tlen > TITLE_MAX * 2:
        fail(f"p{idx}: Title length {tlen:.0f} chars (> {TITLE_MAX*2}; cannot fit 2 lines; sharpen assertion): 「{t}」")
    elif tlen > TITLE_MAX and not explicit_break:
        warn(f"p{idx}: Title length {tlen:.0f} chars (Will span 2 lines; ensure clean break and no orphans): 「{t}」")
    if re.match(r"^(Step|STEP|ステップ)\s*\d", t):
        fail(f"p{idx}: Step tag concatenated in title (use kicker tag chip instead): 「{t}」")
    if re.search(r"^(この|その|ここまで|This|That|These|Those)", t):
        fail(f"p{idx}: Title begins with referential pronoun: 「{t}」")
    if re.match(r"^(まずは|では|そして|さらに|ちなみに)|^(まず|また|次に)[、,]|^(Firstly|Secondly|Furthermore|Moreover|In addition)[, ]", t):
        warn(f"p{idx}: Title begins with spoken transition connector / 接続詞で始まる (Start with subject -- slide-rules §2.18): 「{t}」")
    if t.count("（") + t.count("(") >= 2:
        warn(f"p{idx}: Multiple parenthetical phrases in title: 「{t}」")
    if not TEMPLATE_MODE and re.search(r"[◯○]{2,}|Text\s*\d|ラベル\s*\d|Label\s*\d|タイトル\s*\d|Title\s*\d|Source\s*\d|YYYY|ダミー|Dummy|^資料名$|^会社名$|^Deck Title$|^Company Name$", t):
        fail(f"p{idx}: Template placeholder remains in title / プレースホルダー (Write titles from storyline -- slide-rules §2.8): 「{t}」")


def check_title_variety(titles):
    real = [t.strip() for t in titles if t and t.strip()]
    if len(real) < 5:
        return
    molds = {
        "「◯◯は、…」 / 'X delivers Y'": lambda t: re.match(r"^.{1,12}は[、,]", t),
        "Numbers as subject ('3 key pillars...')": lambda t: re.match(r"^[0-9０-９一二三四五六七八九十]+(?:つ|つの| steps| pillars)", t, re.I),
    }
    for name, f in molds.items():
        hits = [t for t in real if f(t)]
        if len(hits) >= max(4, int(len(real) * 0.6)):
            warn(f"Title sentence structures skewed toward {name} ({len(hits)}/{len(real)} slides). Vary structures to fit individual messages -- slide-rules §2.8")


# --- Count mismatches (slide-rules §2.9) ---------------------------------------------
KANSUJI = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
           "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
COUNT_WORDS = r"(?:段階|フェーズ|ステップ|柱|論点|phases?|steps?|pillars?)"
ENUM_FAMILY = [
    (r"段階|フェーズ|ステップ|phases?|steps?", r"(?:フェーズ|ステップ|STEP|Step|PHASE|Phase|段階)"),
    (r"柱|pillars?", r"(?:柱|Pillar|PILLAR)"),
    (r"論点|issues?", r"(?:論点|Issue|ISSUE)"),
]


def _to_int(s):
    s = s.translate(str.maketrans("０１２３４５６７８９", "0123456789"))
    if s.isdigit():
        return int(s)
    return KANSUJI.get(s)


def check_count_match(idx, title, text):
    if not title or not text:
        return
    for m in re.finditer(r"([0-9０-９]+|[一二三四五六七八九十])\s*(" + COUNT_WORDS + ")", title, re.I):
        n = _to_int(m.group(1))
        if not n or not (2 <= n <= 12):
            continue
        fam = next((lab for pat, lab in ENUM_FAMILY if re.search(pat, m.group(2), re.I)), None)
        if not fam:
            continue
        found = [v for v in (_to_int(x.group(1))
                             for x in re.finditer(fam + r"\s*([0-9０-９]+)", text, re.I)) if v]
        if found and max(found) != n:
            fail(f"p{idx}: Count in title 「{m.group(0)}」 mismatches body numbering (max {max(found)}) (§2.9): 「{title}」")


# ---------------------------------------------------------------- PPTX
def _fwlen(t):
    return sum(0.5 if ord(ch) < 0x3000 else 1 for ch in t)


def check_orphan(idx, shape, label="タイトル"):
    try:
        if not (shape.has_text_frame and shape.width):
            return
        tf = shape.text_frame
        if tf.word_wrap is False:
            return
        raw = tf.text
        if not raw.strip():
            return
        szs = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
        if not szs:
            return
        pt = max(szs)
        w_in = shape.width / 914400.0
        cap = max(4, int(w_in / (pt / 72.0)))
        for line in raw.replace("\v", "\n").split("\n"):
            L = _fwlen(line.strip())
            if L <= cap:
                continue
            last = L - cap * ((int(-(-L // cap))) - 1)
            if 0 < last <= 3:
                warn(f"p{idx}: {label} may have orphan wrap / 泣き別れ (estimated line {int(-(-L // cap))} has only {last:.0f} chars; break at semantic boundary -- §2.13): 「{line.strip()[:40]}」")
    except Exception:
        return


def check_fill_ratio(idx, slide, title_shape):
    try:
        top_lim = 1.6 * 914400
        bot_lim = 6.8 * 914400
        tops, bots = [], []
        for sh in slide.shapes:
            if sh is title_shape or sh.top is None or sh.height is None:
                continue
            t, b = sh.top, sh.top + sh.height
            if b <= top_lim or t >= bot_lim:
                continue
            if sh.height < 914400 * 0.02 and sh.width and sh.width > 914400 * 11:
                continue
            tops.append(max(t, top_lim)); bots.append(min(b, bot_lim))
        if not tops:
            return
        cov = (max(bots) - min(tops)) / (bot_lim - top_lim)
        if cov < 0.55:
            warn(f"p{idx}: Canvas fill ratio {cov*100:.0f}% (body occupies less than 55% of canvas; add substance or switch archetype -- §5.13)")
    except Exception:
        return


def check_pptx(path):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("python-pptx is required: pip3 install python-pptx")
    from pptx_open import open_presentation
    prs = open_presentation(path)
    want_w, want_h = (HOUSE["slide"]["emu"] if HOUSE else (EMU_W, EMU_H))
    if not XML_ONLY and (abs(prs.slide_width - want_w) > 2000 or abs(prs.slide_height - want_h) > 2000):
        fail(f"Slide dimensions {prs.slide_width}x{prs.slide_height} ≠ {'House' if HOUSE else '16:9'} {want_w}x{want_h}")

    titles = []
    term_pages = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        for n in sorted(names, key=lambda s: int(re.search(r"slide(\d+)", s).group(1))):
            xml = z.read(n).decode("utf8", "ignore")
            idx = int(re.search(r"slide(\d+)", n).group(1))
            term_pages.append((idx, pptx_text(xml)))
            check_xml_validity(idx, xml)
            if re.search(r'<p:sld\b[^>]*\bshow="0"', xml):
                warn(f"p{idx}: Hidden slide detected (delete if unused; permitted if intentional appendix backup -- slide-rules §8)")
            if XML_ONLY:
                continue
            rr = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                if 'prst="roundRect"' in sp:
                    h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', sp)
                    if h and int(h.group(1)) >= 365760:
                        rr += 1
            if rr:
                fail(f"p{idx}: Large rounded rectangle / roundRect detected ×{rr} (use sharp right-angle rectangles)")
            fb = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                spPr = re.search(r"<p:spPr>.*?</p:spPr>", sp, re.S)
                if not spPr:
                    continue
                pr = spPr.group(0)
                h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', pr)
                if not (h and int(h.group(1)) >= 365760):
                    continue
                ln = re.search(r"<a:ln[ >].*?</a:ln>", pr, re.S)
                filled = "<a:solidFill>" in re.sub(r"<a:ln[ >].*?</a:ln>", "", pr, flags=re.S)
                if filled and ln and "<a:solidFill>" in ln.group(0):
                    fb += 1
            if fb:
                warn(f"p{idx}: Border on filled shape / 塗りあり図形に枠線 ×{fb} (Filled containers should have no borders -- slide-rules §5.3)")
            for c in OLD_COLORS:
                if f'val="{c}"' in xml or f'val="{c.lower()}"' in xml:
                    fail(f"p{idx}: Prohibited color #{c}")
            lb = len(re.findall(r"<a:t>\s*•", xml))
            if lb:
                fail(f"p{idx}: Direct bullet character '•' hardcoded in text ×{lb} (use native buChar formatting -- slide-rules §7.3)")
            ld = len(re.findall(r"<a:t>\s*[–‐-]\s\s", xml))
            if ld:
                warn(f"p{idx}: Direct level-2 bullet '– ' hardcoded in text ×{ld} (use native buChar formatting -- slide-rules §7.3)")
        for n in z.namelist():
            if "theme" in n or "slideMaster" in n or "slideLayout" in n:
                xml = z.read(n).decode("utf8", "ignore")
                for c in OLD_COLORS:
                    if f'val="{c}"' in xml:
                        warn(f"{n}: Prohibited color #{c}")

    if XML_ONLY:
        return titles
    if HOUSE:
        with zipfile.ZipFile(path) as z:
            for i, s in enumerate(prs.slides, 1):
                check_house(i, s, z.read(s.part.partname.lstrip("/")).decode("utf8", "ignore"))
    for i, s in enumerate(prs.slides, 1):
        title = ""
        named = next((sh for sh in s.shapes if sh.has_text_frame and sh.name.startswith("Title")), None)
        if s.shapes.title is not None and s.shapes.title.has_text_frame:
            title = s.shapes.title.text_frame.text
        elif named is not None:
            title = named.text_frame.text
        else:
            cands = []
            for sh in s.shapes:
                if sh.has_text_frame and sh.top is not None and sh.top < Emu(1097280):
                    sz = max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)
                    cands.append((sz, sh.text_frame.text))
            if cands:
                title = max(cands)[1]
        title_shape = None
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text == title:
                title_shape = sh
                break
        if title_shape is not None:
            check_orphan(i, title_shape, "Title / タイトル")
        if i == 1:
            for sh in s.shapes:
                if sh is not title_shape and sh.has_text_frame:
                    szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    if szs and max(szs) >= 13 and len(sh.text_frame.text) < 80:
                        check_orphan(i, sh, "Cover text / 表紙の文字")
        else:
            check_fill_ratio(i, s, title_shape)
        explicit_break = "\n" in title.strip() and all(_fwlen(l.strip()) <= TITLE_MAX for l in title.split("\n"))
        title = title.replace("\n", " ")
        titles.append(title)
        check_title(i, title, explicit_break)
        check_count_match(i, title, dict(term_pages).get(i, ""))
        for sh in s.shapes:
            if sh.has_text_frame and sh.top is not None and Emu(1097280) <= sh.top < Emu(1600200):
                txt = sh.text_frame.text.strip()
                if txt and "\n" not in txt and len(txt) < 60 and txt != title:
                    szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    if szs and max(szs) <= 12 and not any(r.font.bold for p in sh.text_frame.paragraphs for r in p.runs):
                        warn(f"p{i}: Suspected subtitle line directly beneath title: 「{txt}」")
    check_terms(term_pages)
    check_number_consistency(term_pages)
    check_ai_smell(term_pages)
    return titles


# ---------------------------------------------------------------- HTML
def check_kicker_and_conclusion(html):
    body = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.S)
    texts = [re.sub(r"\s+", " ", t).strip() for t in re.findall(r">([^<>]{3,60})<", body)]
    kick = [t for t in texts if re.fullmatch(r"(?:\d{2}\s*[・·/|]\s*)?[A-Z][A-Z0-9 &/·・\-]{5,}", t)
            and re.search(r"[A-Z]{3,}\s+[A-Z]{2,}", t)
            and not re.search(r"(?i)confidential|appendix|section|step|page", t)]
    if kick:
        warn(f"Uppercase decorative kicker detected ×{len(kick)}: {' / '.join(sorted(set(kick))[:4])} (§7.22)")
    heads = [re.sub(r"<[^>]+>", "", t).strip() for t in re.findall(r"<(?:h3|h4|div class=\"(?:hd|colhd|colh|col-h)[^\"]*\")[^>]*>(.*?)</", body, re.S)]
    dakara = [t for t in heads if re.match(r"^(だから|なので|つまり|Consequently|Therefore|Thus)[、,:： ]?", t, re.I)]
    if dakara:
        warn(f"Column header begins with conjunction / 左右カラムの見出しが接続詞で始まる ×{len(dakara)} (§4.49: Noun phrases only)")


def _leaf_texts(fragment):
    frag = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", fragment, flags=re.S)
    frag = re.sub(r"<!--.*?-->", " ", frag, flags=re.S)
    frag = re.sub(r"<br\s*/?>", "\n", frag)
    blocks = re.split(r"</?(?:div|p|li|td|th|h[1-6]|section|ul|ol|tr|table)\b[^>]*>", frag)
    out = []
    for b in blocks:
        for piece in re.split(r"\n|•", re.sub(r"<[^>]+>", "", b)):
            piece = re.sub(r"\s+", " ", piece).strip()
            if piece:
                out.append(piece)
    return out


def check_exec_summary(idx, title, fragment):
    if not re.search(r"エグゼクティブサマリー|エグゼクティブ・サマリー|Executive Summary", title, re.I):
        return
    if not re.search(r"<li\b(?:(?!</li>).)*?<(?:ul|ol)\b", fragment, re.S):
        warn(f"p{idx}: Executive summary lacks nested bullets / エグゼクティブサマリーに入れ子のブレットが無い (§7.16: format as primary takeaway + 2-3 indented details)")


def check_html(path):
    global STRICT_LEN
    STRICT_LEN = False
    html = Path(path).read_text(encoding="utf8", errors="ignore")
    if not ("338.67mm" in html and "190.5mm" in html):
        fail("16:9 canvas dimensions (338.67mm x 190.5mm) not found in CSS")
    if re.search(r"@page\s*{[^}]*297mm", html):
        fail("@page retains legacy A4 landscape width (297mm)")
    for c in OLD_COLORS:
        if re.search(c, html, re.I):
            fail(f"Prohibited color #{c}")
    for m in re.finditer(r"([^{}]{0,80}){[^}]*?border-radius\s*:\s*(\d+(?:\.\d+)?)(px|mm|rem|em)", html):
        sel, v, u = m.group(1), float(m.group(2)), m.group(3)
        if re.search(r"pill|chip|tag|badge|dot", sel, re.I):
            continue
        px = v * {"px": 1, "mm": 3.78, "rem": 16, "em": 16}[u]
        if px >= 4:
            fail(f"border-radius {v}{u} on `{sel.strip()[-40:]}` (No rounded corners; sharp right angles only)")
            break
    for cls in ["figttl", "subtitle"]:
        if re.search(r'class="[^"]*\b' + cls + r'\b', html):
            fail(f"Legacy subtitle/figure-label class `.{cls}` remains")
    for cls in ["sub", "lead", "caption"]:
        n = len(re.findall(r'class="[^"]*\b' + cls + r'\b', html))
        if n:
            warn(f"`.{cls}` ×{n} -- Permitted on title page; prohibited beneath titles on content slides")
    if re.search(r'<span class="ac">', html):
        fail("Title color split <span class=\"ac\"> remains")
    th = re.search(r"\bth\s*{[^}]*font-size\s*:\s*(\d+)px", html)
    td = re.search(r"\btd\s*{[^}]*font-size\s*:\s*(\d+)px", html)
    if th and td and int(th.group(1)) < int(td.group(1)) + 2:
        fail(f"Table header {th.group(1)}px is smaller than body {td.group(1)}px + 2pt")
    elif th and td and int(th.group(1)) < int(td.group(1)) + 3:
        warn(f"Table header {th.group(1)}px is less than body + 3px (§6: target header at +3~4pt above body)")
    if re.search(r"\bth\s*{[^}]*color\s*:\s*#?(9[0-9a-f]{5}|a[0-9a-f]{5}|b[0-9a-f]{5}|c[0-9a-f]{5}|888|999|aaa|bbb|ccc|gr[ae]y)\b", html, re.I):
        warn("Table header color is muted gray (§6: header must be primary dark ink)")
    check_kicker_and_conclusion(html)
    if re.search(r"\.q3\s+i\s*{[^}]*clip-path\s*:\s*polygon\(\s*50%\s+0\s*,", html):
        fail("3/4 Harvey ball missing center vertex (50% 50%) / ハーベイボール (polygon must begin with 50% 50%)")
    if not re.search(r"<meta[^>]+charset\s*=\s*[\"']?utf-?8", html, re.I):
        fail('<meta charset="utf-8"> missing / charset (Characters will corrupt on Windows browsers -- slide-rules §8)')
    if re.search(r"\bth\s*{[^}]*font-weight\s*:\s*(400|normal|300)", html):
        fail("Table header has normal/thin font weight (must be bold)")
    if re.search(r"tr:nth-child\((even|odd)\)", html):
        fail("Zebra striping detected in table rules")
    if re.search(r"\btd\b[^{}]*{[^}]*border-bottom\s*:", html) and not re.search(
            r"tr:last-child[^{}]*{[^}]*border(-bottom)?\s*:\s*(0|none)", html):
        fail("Final row bottom border not eliminated (`tr:last-child td{border-bottom:0}` required -- slide-rules §5.4)")
    body_html = html.split("</style>", 1)[1] if "</style>" in html else html
    used_classes = set(re.findall(r'class="([^"]*)"', body_html))
    used_tokens = set(tok for cl in used_classes for tok in cl.split())
    for m in re.finditer(r"([^{}]{0,80}){([^}]*)}", html):
        sel, body = m.group(1), m.group(2)
        if re.search(r"pill|chip|tag|badge|dot|legend", sel, re.I):
            continue
        cls_in_sel = re.findall(r"\.([A-Za-z0-9_-]+)", sel)
        if cls_in_sel and cls_in_sel[-1] not in used_tokens:
            continue
        has_fill = re.search(r"background(-color)?\s*:\s*(?!none|transparent)#?\w", body)
        has_border = re.search(r"border\s*:\s*(?!0|none)\d", body)
        if has_fill and has_border and re.search(r"card|box|pillar|mem|step|stat", sel, re.I):
            warn(f"Border on filled container: `{sel.strip()[-40:]}` (Filled containers must have border:0 -- slide-rules §5.3)")
    titles = []
    pat = r'<(?:section|div)[^>]*class="(?:[^"]*\bslide\b[^"]*|s|s [^"]*)"[^>]*>'
    parts = re.split(pat, html)
    slides = parts[1:] if len(parts) > 1 else []
    slide_classes = re.findall(pat.replace('class="(?:', 'class="((?:', 1).replace(')"[^>]*>', '))"[^>]*>', 1), html)
    for i, s in enumerate(slides, 1):
        m = re.search(r"<h1[^>]*>(.*?)</h1>", s, re.S) or re.search(r'class="[^"]*\b(?:ttl|title|msg)\b[^"]*"[^>]*>(.*?)</', s, re.S)
        t = re.sub(r"<[^>]+>", "", m.group(1)).strip() if m else ""
        t = re.sub(r"\s+", " ", t)
        titles.append(t)
        check_title(i, t)
        check_count_match(i, t, re.sub(r"<[^>]+>", " ", s))
        leaves = _leaf_texts(s)
        check_body_placeholders(i, leaves, t)
        check_exec_summary(i, t, s)
        if "cover" not in (slide_classes[i - 1] if i - 1 < len(slide_classes) else ""):
            check_multi_sentence(i, leaves)
    if not slides:
        warn("No `.slide` elements found (title check skipped)")
    check_production_meta(slides if slides else [html])
    if slides:
        check_terms([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)])
        check_number_consistency([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)])
        check_ai_smell([(i, re.sub(r"<[^>]+>", "\n", s)) for i, s in enumerate(slides, 1)])
        check_forbidden([(i, re.sub(r"<[^>]+>", " ", s)) for i, s in enumerate(slides, 1)], html)
    else:
        check_terms([(1, re.sub(r"<[^>]+>", " ", html))])
        check_ai_smell([(1, re.sub(r"<[^>]+>", "\n", html))])
    return titles


def main():
    global TEMPLATE_MODE
    args = sys.argv[1:]
    global FORBIDDEN_TERMS, PART_TYPE_NAMES, XML_ONLY, HOUSE
    if "--xml-only" in args:
        XML_ONLY = True
        args.remove("--xml-only")
    if "--house" in args:
        k = args.index("--house")
        import json
        HOUSE = json.loads(Path(args[k + 1]).read_text(encoding="utf8"))
        del args[k:k + 2]
    if "--template" in args:
        TEMPLATE_MODE = True
        args.remove("--template")
    if "--forbid" in args:
        k = args.index("--forbid")
        FORBIDDEN_TERMS = load_forbidden(args[k + 1])
        del args[k:k + 2]
    PART_TYPE_NAMES = _part_type_names()
    if not args:
        sys.exit(__doc__)
    p = args[0]
    titles = check_pptx(p) if p.lower().endswith((".pptx", ".potx")) else check_html(p)
    if not TEMPLATE_MODE:
        check_title_variety(titles)
    print("=== Slide Titles (Read sequentially to verify narrative storyline) / タイトル一覧 ===")
    for i, t in enumerate(titles, 1):
        print(f"{i:>3}  {t or '(None)'}")
    print()
    for w in warns:
        print("WARN ", w)
    for f in fails:
        print("FAIL ", f)
    print(f"\n{len(fails)} FAIL / {len(warns)} WARN")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()

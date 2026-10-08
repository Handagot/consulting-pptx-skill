#!/usr/bin/env python3
"""Mechanically detect layout corruptions in pptx files.

Detects:
  hidden   ... text covered by opaque shapes drawn on top, making it unreadable
  empty    ... empty boxes with neither text nor images (forgotten callouts, borders)
  outside  ... elements extending beyond the slide boundary
  overflow ... text volume too large for its bounding box
  squash   ... image aspect ratio significantly distorted compared to original

Usage: python3 scripts/check_deck_layout.py <pptx> [page_numbers...]

All results are treated as WARN (exit code is always 0).
Overflow is estimated from character count and may differ from actual rendering.
Only flags empty standalone boxes; blank slides and empty tables are out of scope.
HTML is out of scope (use check_layout.mjs for HTML).
"""
import math
import sys

if hasattr(sys.stdout, "reconfigure") and sys.stdout.encoding:
    sys.stdout.reconfigure(errors="replace")
if hasattr(sys.stderr, "reconfigure") and sys.stderr.encoding:
    sys.stderr.reconfigure(errors="replace")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

EMU_IN = 914400


def walk(shapes, dx=0, dy=0, sx=1.0, sy=1.0):
    for sh in shapes:
        try:
            l = dx + sh.left * sx
            t = dy + sh.top * sy
            w = sh.width * sx
            h = sh.height * sy
        except TypeError:
            continue
        yield sh, (l, t, w, h)
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            g = sh._element.find(".//" + qn("a:xfrm"))
            cx = cy = 0
            cw, ch = sh.width, sh.height
            if g is not None:
                co, ce = g.find(qn("a:chOff")), g.find(qn("a:chExt"))
                if co is not None:
                    cx, cy = int(co.get("x")), int(co.get("y"))
                if ce is not None:
                    cw = int(ce.get("cx")) or sh.width
                    ch = int(ce.get("cy")) or sh.height
            rx = (sh.width / cw) if cw else 1
            ry = (sh.height / ch) if ch else 1
            yield from walk(sh.shapes, dx + (sh.left - cx * rx) * sx,
                            dy + (sh.top - cy * ry) * sy, sx * rx, sy * ry)


def overlap(a, b):
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    ix = max(0, min(ax + aw, bx + bw) - max(ax, bx))
    iy = max(0, min(ay + ah, by + bh) - max(ay, by))
    return ix * iy


def is_opaque(sh):
    """Check whether shape obscures what lies underneath via fill or image"""
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    try:
        f = sh.fill
        if f.type is None:
            return False
        return "SOLID" in str(f.type) or "PATTERN" in str(f.type) or "PICTURE" in str(f.type)
    except Exception:
        return False


def text_of(sh):
    tf = getattr(sh, "text_frame", None)
    return tf.text.strip() if tf is not None else ""


def font_pt(sh, default=10.0):
    tf = getattr(sh, "text_frame", None)
    if tf is None:
        return default
    for p in tf.paragraphs:
        for r in p.runs:
            if r.font.size:
                return r.font.size.pt
    return default


def check(path, pages=None):
    prs = Presentation(path)
    SW, SH = prs.slide_width, prs.slide_height
    issues = []
    for n, slide in enumerate(prs.slides, 1):
        if pages and n not in pages:
            continue
        items = list(walk(slide.shapes))
        for i, (sh, box) in enumerate(items):
            t = text_of(sh)
            l, top, w, h = box
            if w <= 0 or h <= 0:
                continue
            # Check if outside boundary
            if l < -w * 0.5 or top < -h * 0.5 or l + w > SW + w * 0.5 or top + h > SH * 1.02:
                issues.append((n, "outside", f"{(t or sh.name)[:24]} is out of bounds ({l/EMU_IN:.1f}in, {top/EMU_IN:.1f}in)"))
            if t:
                # Check if covered by opaque shape on top
                covered = 0
                for sh2, box2 in items[i + 1:]:
                    if sh2 is sh or not is_opaque(sh2):
                        continue
                    if text_of(sh2) and overlap(box, box2) < w * h * 0.9:
                        continue
                    covered = max(covered, overlap(box, box2))
                if covered > w * h * 0.85:
                    issues.append((n, "hidden", f"'{t[:26]}' is obscured by another shape"))
                # Check text overflow (rough estimate)
                pt = font_pt(sh)
                cpl = max(1, int((w / EMU_IN * 72) / (pt * 1.02)))
                lines = sum(max(1, math.ceil(len(p.text) / cpl)) for p in sh.text_frame.paragraphs)
                need = lines * pt * 1.45 / 72 * EMU_IN
                if top + need > SH:
                    issues.append((n, "outside", f"'{t[:22]}' overflows outside the slide"))
                elif need > h * 1.3 and h > 100000:
                    issues.append((n, "overflow", f"'{t[:22]}' overflows its box (needs {need/EMU_IN:.1f}in / box {h/EMU_IN:.1f}in)"))
            else:
                # Boxes with neither text nor image (border-only highlights are allowed, only check filled shapes)
                # Shapes with height or width < 0.05in are lines/dividers, not treated as boxes
                thin = min(w, h) < 0.05 * EMU_IN
                if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and is_opaque(sh) and not thin:
                    # Highlight frames over pictures or background panels under text are acceptable
                    on_pic = any(sh2.shape_type == MSO_SHAPE_TYPE.PICTURE and overlap(box, box2) > w * h * 0.5
                                 for sh2, box2 in items)
                    has_label = any(text_of(sh2) and overlap(box, box2) > w * h * 0.3
                                    for sh2, box2 in items[i + 1:])
                    if not on_pic and not has_label:
                        issues.append((n, "empty", f"Empty box without content ({w/EMU_IN:.1f}x{h/EMU_IN:.1f}in)"))
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                # Check if image is covered by filled shape
                for sh2, box2 in items[i + 1:]:
                    if sh2.shape_type == MSO_SHAPE_TYPE.PICTURE or not is_opaque(sh2):
                        continue
                    if text_of(sh2):
                        continue
                    if overlap(box, box2) > w * h * 0.15:
                        issues.append((n, "hidden", f"{overlap(box, box2)/(w*h)*100:.0f}% of image covered by filled shape"))
                        break
                try:
                    iw, ih = sh.image.size
                    if iw and ih:
                        r = (w / h) / (iw / ih)
                        if r > 1.25 or r < 0.8:
                            issues.append((n, "squash", f"Image aspect ratio distorted by {r:.2f}x"))
                except Exception:
                    pass
    return issues


if __name__ == "__main__":
    path = sys.argv[1]
    pages = [int(x) for x in sys.argv[2:]] or None
    issues = check(path, pages)
    for n, kind, msg in issues:
        print(f"p{n:<3} [{kind}] {msg}")
    print(f"--- {len(issues)} issues")
